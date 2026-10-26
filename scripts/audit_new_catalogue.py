import openpyxl
import json
from collections import Counter

file_path = "al burhani new stock 27-10-2026.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb["in"]

headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
print("Headers:", headers)

rows_data = []
row_idx = 1
for row in ws.iter_rows(min_row=2, values_only=False):
    row_idx += 1
    vals = [cell.value for cell in row]
    if not any(vals):
        continue
    
    sku_cell = row[0]
    barcode_cell = row[1]
    name_cell = row[2]
    cat_cell = row[3]
    unit_cell = row[4]
    sp_cell = row[5]
    cp_cell = row[6]
    stock_cell = row[7]
    min_stock_cell = row[8]
    status_cell = row[9]

    # Inspect cell types and string formatting
    rows_data.append({
        "row_number": row_idx,
        "sku_raw": sku_cell.value,
        "sku_type": sku_cell.data_type,
        "barcode_raw": barcode_cell.value,
        "barcode_type": barcode_cell.data_type,
        "barcode_fmt": barcode_cell.number_format,
        "name": name_cell.value,
        "category": cat_cell.value,
        "unit": unit_cell.value,
        "selling_price": sp_cell.value,
        "cost_price": cp_cell.value,
        "current_stock": stock_cell.value,
        "min_stock": min_stock_cell.value,
        "status": status_cell.value,
    })

print(f"\nTotal data rows parsed: {len(rows_data)}")

# Validation checks
missing_names = []
missing_barcodes = []
missing_skus = []
duplicate_skus = []
duplicate_barcodes = []
invalid_prices = []
invalid_stocks = []

skus = []
barcodes = []
categories = set()
units = set()
statuses = set()

for r in rows_data:
    # Name check
    name = str(r["name"]).strip() if r["name"] is not None else ""
    if not name:
        missing_names.append(r["row_number"])
    
    # SKU check
    sku_val = r["sku_raw"]
    if sku_val is None or str(sku_val).strip() == "":
        missing_skus.append(r["row_number"])
        sku_str = ""
    else:
        # Check if integer or float
        if isinstance(sku_val, float) and sku_val.is_integer():
            sku_str = str(int(sku_val))
        else:
            sku_str = str(sku_val).strip()
    skus.append((sku_str, r["row_number"]))

    # Barcode check
    bc_val = r["barcode_raw"]
    if bc_val is None or str(bc_val).strip() == "":
        missing_barcodes.append(r["row_number"])
        bc_str = ""
    else:
        if isinstance(bc_val, float):
            # check if integer
            if bc_val.is_integer():
                bc_str = str(int(bc_val))
            else:
                bc_str = f"{bc_val:.0f}"
        elif isinstance(bc_val, int):
            bc_str = str(bc_val)
        else:
            bc_str = str(bc_val).strip()
    barcodes.append((bc_str, r["row_number"]))

    # Category, Unit, Status
    if r["category"] is not None:
        categories.add(str(r["category"]).strip())
    if r["unit"] is not None:
        units.add(str(r["unit"]).strip())
    if r["status"] is not None:
        statuses.add(str(r["status"]).strip())

    # Prices
    sp = r["selling_price"]
    try:
        sp_num = float(sp) if sp is not None else 0.0
        if sp_num < 0:
            invalid_prices.append((r["row_number"], "Negative Selling Price", sp))
    except:
        invalid_prices.append((r["row_number"], "Invalid Selling Price", sp))

    # Stocks
    st = r["current_stock"]
    try:
        st_num = float(st) if st is not None else 0.0
    except:
        invalid_stocks.append((r["row_number"], "Invalid Stock", st))

# Duplicate detection
sku_counts = Counter([s for s, _ in skus if s])
barcode_counts = Counter([b for b, _ in barcodes if b])

dup_skus = {s: count for s, count in sku_counts.items() if count > 1}
dup_bcs = {b: count for b, count in barcode_counts.items() if count > 1}

print("\n--- Summary of Analysis ---")
print(f"Missing names: {len(missing_names)}")
print(f"Missing SKUs: {len(missing_skus)}")
print(f"Missing Barcodes: {len(missing_barcodes)}")
print(f"Duplicate SKUs: {dup_skus}")
print(f"Duplicate Barcodes: {dup_bcs}")
print(f"Invalid Prices: {invalid_prices}")
print(f"Invalid Stocks: {invalid_stocks}")

print("\nCategories found (total %d):" % len(categories))
print(sorted(list(categories)))

print("\nUnits found:")
print(sorted(list(units)))

print("\nStatuses found:")
print(sorted(list(statuses)))

# Check sample barcodes to see if any have leading zeros or scientific notation
print("\nSample 10 Barcodes:")
for b, r_idx in barcodes[:10]:
    print(f"Row {r_idx}: '{b}' (len: {len(b)})")

print("\nSample 10 SKUs:")
for s, r_idx in skus[:10]:
    print(f"Row {r_idx}: '{s}'")
